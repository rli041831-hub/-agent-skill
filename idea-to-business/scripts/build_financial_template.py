"""Create an editable, blank simplified unit-economics/cash template."""
import argparse
from pathlib import Path


def build(output):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.workbook.properties import CalcProperties
    from openpyxl.worksheet.datavalidation import DataValidation

    if output.exists():
        raise FileExistsError(f"will not overwrite {output}")
    wb = Workbook()
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    p = wb.active
    p.title = '参数与情景'
    p.append(['简化单位经济与现金模板：待填参数，不是已验证预测'])
    p.append(['适用前提：每月销量与单价固定，即收即付；税费、融资、赊账、库存与季节性另建模型。'])
    p.append(['参数/结果', '保守', '基准', '较好', '单位/说明', '依据与验证状态'])
    labels = [
        ('月销量', '单位/月；单位可为订单或服务，按项目定义'),
        ('单位净收入', '元/单位；不把交易额当平台收入'),
        ('单位变动成本', '元/单位；含对应直接支出'),
        ('月固定现金支出', '元/月'),
        ('一次启动投入', '元；发生于第0月'),
        ('初始可用现金', '元；不含尚未到账融资'),
        ('安全缓冲', '元；由业务约束设定'),
    ]
    for label, unit in labels:
        p.append([label, None, None, None, unit, '待填写来源/估算方法；未验证'])
    p.cell(12, 1, '以上每个情景的7项参数都填写后才显示结果；合法的0值可填写0。')
    rows = {15:'月收入',16:'月变动成本',17:'月贡献',18:'月经营净流入（即收即付）',
            19:'盈亏平衡销量',20:'12个月最低现金（含启动时点）',21:'维持安全缓冲的额外资金缺口'}
    for r, label in rows.items(): p.cell(r, 1, label)
    for col in 'BCD':
        exprs = {15:f'{col}4*{col}5',16:f'{col}4*{col}6',17:f'{col}15-{col}16',18:f'{col}17-{col}7',
                 19:f'IF({col}5<={col}6,"单位贡献不为正",ROUNDUP({col}7/({col}5-{col}6),0))',
                 20:f'{col}9-{col}8+MIN(0,{col}18)*12',21:f'MAX(0,{col}10-{col}20)'}
        for row, expr in exprs.items(): p[f'{col}{row}'] = f'=IF(COUNT({col}4:{col}10)<7,"待填",{expr})'
    dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1=0, allow_blank=True)
    dv.errorTitle = '参数需要非负数'
    dv.error = '请输入非负数字；未知值留空，不用猜测。'
    dv.showErrorMessage = True
    p.add_data_validation(dv)
    dv.add('B4:D10')
    for row in p.iter_rows(min_row=4, max_row=10, min_col=2, max_col=4):
        for cell in row:
            cell.fill = PatternFill('solid', fgColor='E8F1FF')
            cell.number_format = '0.00'
    cash = wb.create_sheet('月度现金')
    cash.append(['月份', '保守期末现金', '基准期末现金', '较好期末现金'])
    for month in range(13):
        row = month + 2
        cash.cell(row, 1, month)
        for col in 'BCD':
            cash[f'{col}{row}'] = f'=IF(COUNT(\'参数与情景\'!{col}4:{col}10)<7,"待填",\'参数与情景\'!{col}9-\'参数与情景\'!{col}8+\'参数与情景\'!{col}18*A{row})'
    note = wb.create_sheet('使用与限制')
    for line in [
        '这是一份空白模板，未填任何市场数据、价格、用户量或收入预测。',
        '本表只适用简化单位模型。订阅留存、不同计价周期、平台费率、收付款时点、税费、库存与融资需按项目改造。',
        '先为各参数补真实来源或明确估算方法，再做情景；三种情景不是已经验证的未来。',
        '改变参数后在电子表格软件中重新计算；不能以文件能打开代替公式结果验收。',
        '现金包含第0月投入，单位贡献非正时不虚构正数盈亏平衡销量。',
        '初始可用现金只记录已可用金额；资金缺口不代表一定能融资或承诺投资回报。',
    ]: note.append([line])
    for sheet in wb:
        sheet.freeze_panes = 'B4' if sheet == p else 'B2'
        for row in sheet:
            for cell in row:
                cell.font = Font(name='Microsoft YaHei', size=11)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        sheet.column_dimensions['A'].width = 44
        for col in 'BCD': sheet.column_dimensions[col].width = 22
    p.column_dimensions['E'].width = 44
    p.column_dimensions['F'].width = 48
    note.column_dimensions['A'].width = 110
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    try: build(args.output)
    except (OSError, ImportError) as error: parser.exit(2, f'Workbook not created: {error}\n')
    print(args.output.resolve())
